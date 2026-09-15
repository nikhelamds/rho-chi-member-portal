"""Read-only polling of the three chapter workbooks using their existing link access."""
import concurrent.futures,datetime,io,json,os,threading,time,urllib.request,zipfile,warnings
from pathlib import Path
import openpyxl
from data_provider import ACT,CAL,RUSH,normalize

def timestamp():return datetime.datetime.now(datetime.timezone.utc).isoformat()
def serial(value):
 if isinstance(value,datetime.datetime):return (value-datetime.datetime(1899,12,30)).total_seconds()/86400
 if isinstance(value,datetime.time):return (value.hour*3600+value.minute*60+value.second)/86400
 return value

def download(book):
 url=f'https://docs.google.com/spreadsheets/d/{book}/export?format=xlsx'
 with urllib.request.urlopen(urllib.request.Request(url,headers={'User-Agent':'RhoChiPortal/1.0'}),timeout=25) as response:
  payload=response.read(12*1024*1024+1)
 if len(payload)>12*1024*1024 or not payload.startswith(b'PK'):raise ValueError('Workbook download unavailable.')
 with zipfile.ZipFile(io.BytesIO(payload)) as z:
  if sum(i.file_size for i in z.infolist())>100*1024*1024:raise ValueError('Workbook is too large.')
 with warnings.catch_warnings():
  warnings.simplefilter('ignore',UserWarning)
  return openpyxl.load_workbook(io.BytesIO(payload),data_only=True)

def collect(loader=download):
 with concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool:act,rush,cal=list(pool.map(loader,[ACT,RUSH,CAL]))
 raw={}
 classes=next((t for t in act.sheetnames if t.startswith('Classes')),None)
 for key,title,limit in [('directory','Actives Fall 2026',200),('attendance','Fall 2026 Attendance',1000),('policy','Point System',100),('buddies','Buddy of the Month',200),('academics',classes,1000)]:
  if title not in act.sheetnames:raise ValueError('An expected worksheet is missing.')
  raw[key]=[[serial(c) for c in row] for row in act[title].iter_rows(max_row=min(act[title].max_row,limit),values_only=True)]
 if 'Dynamic Calendar' not in cal.sheetnames or 'Sheet1' not in rush.sheetnames:raise ValueError('An expected worksheet is missing.')
 raw['calendar']=[[serial(c) for c in r] for r in cal['Dynamic Calendar'].iter_rows(max_row=min(cal['Dynamic Calendar'].max_row,1002),max_col=8,values_only=True)]
 raw['shifts']=[[serial(c) for c in r] for r in rush['Sheet1'].iter_rows(max_row=min(rush['Sheet1'].max_row,200),max_col=15,values_only=True)]
 raw['links']=[]
 for title in ['Actives Fall 2026','Comment Box']:
  if title not in act.sheetnames:continue
  for row in act[title].iter_rows(max_row=17,max_col=7):
   for c in row:
    if c.hyperlink and c.hyperlink.target and c.hyperlink.target.startswith('https://'):
     raw['links'].append(dict(label=str(c.value).strip(),url=c.hyperlink.target,sheet=title,cell=c.coordinate))
 for row in raw['directory'][8:]:
  for i in (2,4,8):
   if len(row)>i:row[i]=None
 checked=normalize(raw,'')
 if not checked['directory'] or not checked['events']:raise ValueError('Workbook response had no member or calendar records.')
 for workbook in [act,rush,cal]:workbook.close()
 return raw

class SheetsSync:
 def __init__(self,private,interval=60,loader=collect):
  self.private=Path(private);self.interval=max(60,interval);self.loader=loader;self.lock=threading.Lock();self.state_lock=threading.Lock();self.stop_event=threading.Event()
  self.path=self.private/'source-data.json';self.state_path=self.private/'sync-status.json'
  try:previous=json.loads(self.state_path.read_text())
  except (OSError,ValueError):previous={}
  self.state=dict(enabled=True,intervalSeconds=self.interval,lastSuccess=previous.get('lastSuccess'),lastAttempt=previous.get('lastAttempt'),error=None,refreshing=False)
 def status(self):
  with self.state_lock:return dict(self.state)
 def atomic(self,path,value):
  temp=path.with_suffix('.tmp');temp.write_text(json.dumps(value));os.replace(temp,path)
 def refresh(self):
  with self.lock:
   with self.state_lock:self.state.update(lastAttempt=timestamp(),refreshing=True)
   try:
    raw=self.loader();raw['_sync']=dict(lastSuccess=timestamp())
    self.atomic(self.path,raw)
    with self.state_lock:self.state.update(lastSuccess=raw['_sync']['lastSuccess'],error=None)
    success=True
   except Exception:
    with self.state_lock:self.state['error']='Could not refresh Google Sheets. Showing the last successful import; retrying automatically.'
    success=False
   finally:
    with self.state_lock:self.state['refreshing']=False
    self.atomic(self.state_path,self.status())
   return success
 def start(self):
  def loop():
   while not self.stop_event.is_set():
    self.refresh()
    self.stop_event.wait(self.interval)
  threading.Thread(target=loop,name='chapter-sheets-sync',daemon=True).start()
